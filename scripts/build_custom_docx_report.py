import os
import re
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

# ---------------------------------------------------------------------------
# Styling and OpenXML Helpers
# ---------------------------------------------------------------------------

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Set inner margins (padding) of a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def prevent_row_split(row):
    """Prevent table row from splitting across pages."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_repeat_header(row):
    """Repeat header row on each page."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Set clean, subtle borders on a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_inline_markdown(p, text, base_font="Times New Roman", base_size=13.0, is_bold_base=False, is_italic_base=False, color=None):
    """Parse inline bold, italic, and code markdown tags and create formatted runs."""
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = is_bold_base
        is_italic = is_italic_base
        is_code = False
        run_text = token

        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            is_bold = True
            run_text = token[2:-2]
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            is_italic = True
            run_text = token[1:-1]
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            is_code = True
            run_text = token[1:-1]

        # Clean LaTeX math delimiters $...$ to plain unicode representation
        run_text = run_text.replace('$', '').replace(r'\le', '≤').replace(r'\ge', '≥')
        run_text = run_text.replace(r'\approx', '≈').replace(r'\to', '→').replace(r'\times', '×')
        run_text = run_text.replace(r'\leftrightarrow', '↔').replace(r'\pm', '±').replace(r'\Delta', 'Δ')
        run_text = run_text.replace(r'\text{ms}', 'ms').replace(r'\text{s}', 's')
        run_text = run_text.replace(r'\text{px}', 'px').replace(r'\text{pixels}', 'pixels').replace(r'\text{fps}', 'fps')
        run_text = run_text.replace(r'\text{KB}', 'KB').replace(r'\text{MB}', 'MB').replace(r'\text{GB}', 'GB')
        run_text = run_text.replace(r'\text{giây}', 'giây').replace(r'\text{E2E}', 'E2E')
        run_text = re.sub(r'\\text\{([^}]+)\}', r'\1', run_text)

        run = p.add_run(run_text)
        f_name = "Consolas" if is_code else base_font
        run.font.name = f_name
        run.font.size = Pt(base_size - 1.0 if is_code else base_size)
        if is_code:
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        else:
            if is_bold:
                run.font.bold = True
            if is_italic:
                run.font.italic = True
            if color:
                run.font.color.rgb = color

        # Hardened Font configuration for Vietnamese Unicode compatibility
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn('w:ascii'), f_name)
        rFonts.set(qn('w:hAnsi'), f_name)
        rFonts.set(qn('w:eastAsia'), f_name)
        rFonts.set(qn('w:cs'), f_name)

def add_callout_box(doc, text_lines, title="ĐIỂM NHẤN KIẾN TRÚC & NGUYÊN TẮC THIẾT KẾ"):
    """Create a stunning modern callout box with a colored left-accent border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.2)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=160)
    
    # Left border accent bar (3pt solid #2563EB)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    # Title of Callout
    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run(f"📌 {title}")
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(11.5)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    # Content of Callout
    for line in text_lines:
        p_c = cell.add_paragraph()
        p_c.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_c.paragraph_format.space_before = Pt(2)
        p_c.paragraph_format.space_after = Pt(4)
        p_c.paragraph_format.line_spacing = 1.3
        format_inline_markdown(p_c, line, base_size=11.5, is_italic_base=True, color=RGBColor(0x33, 0x41, 0x55))
        
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ---------------------------------------------------------------------------
# Main Document Builder
# ---------------------------------------------------------------------------

def build_report():
    input_md_path = 'docs/vi/reports/BAO_CAO_TONG_HOP_RUT_GON.md'
    output_docx_path = 'docs/vi/reports/BAO_CAO_INKSHIFT_FINAL.docx'
    
    if not os.path.exists(input_md_path):
        print(f"Error: {input_md_path} not found.")
        return

    print("Creating fresh, creative, template-free Word document...")
    doc = docx.Document()

    # 1. Page Geometry (A4 Standards per Nghị định 30/2020/NĐ-CP)
    for section in doc.sections:
        section.page_width = Inches(8.27)   # 210 mm
        section.page_height = Inches(11.69) # 297 mm
        section.top_margin = Inches(0.79)    # 2.0 cm
        section.bottom_margin = Inches(0.79) # 2.0 cm
        section.left_margin = Inches(1.18)   # 3.0 cm (binding margin)
        section.right_margin = Inches(0.79)  # 2.0 cm
        section.different_first_page_header_footer = True
        
        # Header (pages after cover)
        hdr = section.header
        p_hdr = hdr.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.paragraph_format.space_after = Pt(4)
        r_hdr = p_hdr.add_run("Báo cáo Kỹ thuật Đồ án — InkShift In-Browser Manga Translation")
        r_hdr.font.name = "Times New Roman"
        r_hdr.font.size = Pt(9.5)
        r_hdr.font.italic = True
        r_hdr.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        
        # Footer (pages after cover)
        ftr = section.footer
        p_ftr = ftr.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ftr.paragraph_format.space_before = Pt(4)
        r_ftr = p_ftr.add_run("Trang ")
        r_ftr.font.name = "Times New Roman"
        r_ftr.font.size = Pt(10.0)
        r_ftr.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        # Page Number Field
        fld = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        p_ftr._p.append(fld)
        r_of = p_ftr.add_run(" / ")
        r_of.font.name = "Times New Roman"
        r_of.font.size = Pt(10.0)
        r_of.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        fld_num = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
        p_ftr._p.append(fld_num)

    # 2. Configure & Purge Heading Styles (Gotcha 1)
    for s in doc.styles:
        if s.name.startswith('Heading') or s.name in ['Subtitle', 'TOC Heading']:
            for numPr in s._element.xpath('.//w:numPr'):
                numPr.getparent().remove(numPr)
            for ind in s._element.xpath('.//w:ind'):
                ind.getparent().remove(ind)

    # =======================================================================
    # 3. Creative, Template-Free Cover Page (Bìa Sáng Tạo Chuẩn Cao Cấp)
    # =======================================================================
    print("Designing bespoke creative cover page...")
    
    # Institution & Category Top Tagline
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(20)
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("BÁO CÁO KỸ THUẬT TỔNG HỢP — ĐỒ ÁN MÔN HỌC")
    r_inst.font.name = "Times New Roman"
    r_inst.font.size = Pt(12.0)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_sub_inst = doc.add_paragraph()
    p_sub_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub_inst.paragraph_format.space_before = Pt(0)
    p_sub_inst.paragraph_format.space_after = Pt(12)
    r_sub_inst = p_sub_inst.add_run("CHUYÊN NGÀNH: KỸ THUẬT PHẦN MỀM & TRÍ TUỆ NHÂN TẠO ỨNG DỤNG")
    r_sub_inst.font.name = "Times New Roman"
    r_sub_inst.font.size = Pt(10.5)
    r_sub_inst.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # Decorative Horizontal Accent Line
    p_rule = doc.add_paragraph()
    p_rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rule.paragraph_format.space_before = Pt(0)
    p_rule.paragraph_format.space_after = Pt(45)
    r_rule = p_rule.add_run("━" * 45)
    r_rule.font.name = "Times New Roman"
    r_rule.font.size = Pt(12.0)
    r_rule.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    # Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("INKSHIFT")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(32.0)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Project Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(25)
    r_sub = p_sub.add_run("TIỆN ÍCH DỊCH TRUYỆN TRANH THỜI GIAN THỰC ĐA NỀN TẢNG\nTRÊN TRÌNH DUYỆT WEB")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(14.0)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    # Creative Tech Badges Paragraph
    p_badges = doc.add_paragraph()
    p_badges.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_badges.paragraph_format.space_before = Pt(0)
    p_badges.paragraph_format.space_after = Pt(45)
    r_badges = p_badges.add_run("DirectML FP16  •  YOLO11s-seg  •  RapidOCR PP-OCRv4  •  Shadow DOM v1  •  SQLite WAL")
    r_badges.font.name = "Consolas"
    r_badges.font.size = Pt(9.5)
    r_badges.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Metadata Box / Card (Khối Thông Tin Đề Tài & Nhóm Thực Hiện)
    meta_tbl = doc.add_table(rows=5, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False
    
    col_widths = [Inches(2.2), Inches(4.0)]
    meta_data = [
        ("Cơ quan / Đơn vị đào tạo:", "Khoa Công Nghệ Thông Tin / Bộ Môn Kỹ Thuật Phần Mềm"),
        ("Giảng viên hướng dẫn:", "TS. [Điền tên Giảng viên hướng dẫn]"),
        ("Nhóm sinh viên thực hiện:", "Nhóm InkShift (3 thành viên — AI Lead, Frontend Lead, Gateway Lead)"),
        ("Phiên bản mã nguồn:", "Release 1.0.0 (Phase 32 — RapidOCR DirectML & 1-Click Launch)"),
        ("Môi trường thực thi:", "Windows 11, Chromium MV3, Python 3.12, DirectML ONNX Runtime")
    ]
    
    for row_idx, (label, val) in enumerate(meta_data):
        row = meta_tbl.rows[row_idx]
        prevent_row_split(row)
        
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = col_widths[0]
        c1.width = col_widths[1]
        
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, top=80, bottom=80, left=120, right=80)
        set_cell_margins(c1, top=80, bottom=80, left=80, right=120)
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(11.0)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11.0)
        r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    set_table_borders(meta_tbl, color="CBD5E1", sz="6", val="single")

    # Bottom Location & Date
    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_loc.paragraph_format.space_before = Pt(80)
    p_loc.paragraph_format.space_after = Pt(0)
    r_loc = p_loc.add_run("Hà Nội, Tháng 10 Năm 2026")
    r_loc.font.name = "Times New Roman"
    r_loc.font.size = Pt(12.0)
    r_loc.font.italic = True
    r_loc.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # Page Break after Cover Page
    doc.add_page_break()

    # =======================================================================
    # 4. Parse and Render Markdown Content
    # =======================================================================
    print("Reading and parsing condensed markdown report...")
    with open(input_md_path, 'r', encoding='utf-8') as f:
        md_lines = f.readlines()

    # Start rendering content from Executive Summary onwards
    start_idx = 0
    for i, line in enumerate(md_lines):
        if line.strip().startswith("### TÓM TẮT ĐIỀU HÀNH"):
            start_idx = i
            break

    content_lines = md_lines[start_idx:]
    print(f"Compiling {len(content_lines)} markdown lines...")

    in_table = False
    table_lines = []
    in_code_block = False
    code_lines = []
    first_h1 = True

    def flush_table(tbl_lines):
        if not tbl_lines or len(tbl_lines) < 2:
            return
        parsed_rows = []
        for l in tbl_lines:
            l = l.strip()
            if not l.startswith('|') or not l.endswith('|'):
                continue
            cells = [c.strip() for c in l[1:-1].split('|')]
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            parsed_rows.append(cells)

        if not parsed_rows:
            return

        num_cols = max(len(r) for r in parsed_rows)
        for r in parsed_rows:
            while len(r) < num_cols:
                r.append("")

        tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        set_table_borders(tbl, color="CBD5E1", sz="4", val="single")

        for r_idx, row_data in enumerate(parsed_rows):
            row = tbl.rows[r_idx]
            prevent_row_split(row) # Gotcha 2: cantSplit
            if r_idx == 0:
                set_repeat_header(row) # Gotcha 2: tblHeader

            for c_idx, cell_text in enumerate(row_data):
                cell = row.cells[c_idx]
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

                # Shading & typography
                if r_idx == 0:
                    set_cell_background(cell, "1E293B") # Deep slate header
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "FFFFFF")
                    else:
                        set_cell_background(cell, "F8FAFC") # Soft zebra

                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c_idx == 0 or r_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15

                if r_idx == 0:
                    format_inline_markdown(p, cell_text, base_size=10.5, is_bold_base=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                else:
                    format_inline_markdown(p, cell_text, base_size=10.5, is_bold_base=False, color=RGBColor(0x0F, 0x17, 0x2A))

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def flush_code(c_lines):
        if not c_lines:
            return
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.2)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

        # Subtle left border
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/>'
            f'<w:left w:val="single" w:sz="12" w:space="0" w:color="94A3B8"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1

        full_code = "".join(c_lines).rstrip()
        run = p.add_run(full_code)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn('w:ascii'), "Consolas")
        rFonts.set(qn('w:hAnsi'), "Consolas")

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    i = 0
    while i < len(content_lines):
        line = content_lines[i]
        stripped = line.strip()

        # Handle Code Block fences
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code(code_lines)
                code_lines = []
            else:
                if in_table:
                    in_table = False
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle Markdown Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                flush_table(table_lines)
                table_lines = []

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal rules
        if stripped in ["---", "***", "___"]:
            i += 1
            continue

        # Heading 1 (e.g. ## PHẦN 1:)
        if stripped.startswith("## "):
            h_text = stripped[3:].strip()
            if not first_h1:
                doc.add_page_break()
            else:
                first_h1 = False

            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = p.paragraph_format
            pf.space_before = Pt(18)
            pf.space_after = Pt(10)
            pf.keep_with_next = True # Gotcha 4

            run = p.add_run(h_text.upper())
            run.font.name = "Times New Roman"
            run.font.size = Pt(15.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Royal Navy Accent
            i += 1
            continue

        # Heading 2 (e.g. ### 1.1.)
        if stripped.startswith("### "):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = p.paragraph_format
            pf.space_before = Pt(14)
            pf.space_after = Pt(6)
            pf.keep_with_next = True # Gotcha 4

            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(13.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Slate Dark
            i += 1
            continue

        # Heading 3 (e.g. #### 1.1.1.)
        if stripped.startswith("#### "):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = p.paragraph_format
            pf.space_before = Pt(10)
            pf.space_after = Pt(4)
            pf.keep_with_next = True # Gotcha 4

            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        # Handle Images ![caption](path)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            caption = img_match.group(1)
            rel_path = img_match.group(2)
            # Try to resolve image path
            img_path = rel_path
            if not os.path.exists(img_path):
                img_path = os.path.join('docs/vi/reports', rel_path)
            if not os.path.exists(img_path):
                img_path = os.path.join('docs/vi/reports/images', os.path.basename(rel_path))

            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(10)
                p_img.paragraph_format.space_after = Pt(4)
                p_img.paragraph_format.keep_with_next = True
                
                try:
                    run_img = p_img.add_run()
                    run_img.add_picture(img_path, width=Inches(5.8))
                except Exception as ex:
                    print(f"Warning: could not add picture {img_path}: {ex}")

                # Image Caption (Gotcha 5: Caption ALWAYS BELOW, italic, centered)
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(12)
                format_inline_markdown(p_cap, f"Hình: {caption}", base_size=11.0, is_italic_base=True, color=RGBColor(0x47, 0x55, 0x69))
            else:
                print(f"Warning: image path not found: {rel_path}")
            i += 1
            continue

        # Handle Lists
        list_match = re.match(r'^(\*|-|\d+\.)\s+(.*)', stripped)
        if list_match:
            bullet_sym = list_match.group(1)
            item_text = list_match.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.space_before = Pt(1)
            pf.space_after = Pt(4)
            pf.line_spacing = 1.3
            pf.left_indent = Inches(0.25)

            # Bullet character
            run_b = p.add_run("•  " if bullet_sym in ['*', '-'] else f"{bullet_sym}  ")
            run_b.font.name = "Times New Roman"
            run_b.font.size = Pt(12.5)
            run_b.font.bold = True
            run_b.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

            format_inline_markdown(p, item_text, base_size=12.5)
            i += 1
            continue

        # Standard Body Paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.4
        format_inline_markdown(p, stripped, base_size=12.5)
        i += 1

    # Cleanup table/code if still open at EOF
    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    # =======================================================================
    # 5. Save Document with File-Lock Recovery (Gotcha 7)
    # =======================================================================
    print(f"Saving compiled document to {output_docx_path}...")
    try:
        doc.save(output_docx_path)
        print(f"SUCCESS: Document cleanly compiled and saved to {output_docx_path}")
    except PermissionError:
        fallback_path = 'docs/vi/reports/BAO_CAO_INKSHIFT_FINAL_V2.docx'
        print(f"PermissionError: {output_docx_path} is locked by MS Word! Saving to {fallback_path}...")
        doc.save(fallback_path)
        print(f"SUCCESS: Document saved to fallback path: {fallback_path}")

if __name__ == '__main__':
    build_report()

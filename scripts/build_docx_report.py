import os
import re
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins (padding) of a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def prevent_row_split(row):
    """Prevent table row from splitting across pages."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_repeat_header(row):
    """Repeat header row on each page."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def add_styled_paragraph(doc, text="", style='Normal', space_after=6, space_before=0, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line_spacing
    return p

def format_inline_markdown(p, text, base_font="Times New Roman", base_size=13.0, is_bold_base=False, is_italic_base=False, color=None):
    """Parse inline **bold**, *italic*, and `code` markdown tags and create formatted runs."""
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = is_bold_base
        is_italic = is_italic_base
        is_code = False
        run_text = token

        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            is_bold = True
            run_text = token[2:-2]
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            is_italic = True
            run_text = token[1:-1]
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            is_code = True
            run_text = token[1:-1]

        # Clean LaTeX math delimiters $...$ to plain unicode representation
        run_text = run_text.replace('$', '').replace(r'\le', '≤').replace(r'\ge', '≥')
        run_text = run_text.replace(r'\approx', '≈').replace(r'\to', '→').replace(r'\times', '×')
        run_text = run_text.replace(r'\leftrightarrow', '↔').replace(r'\pm', '±').replace(r'\Delta', 'Δ')
        run_text = run_text.replace(r'\text{ms}', 'ms').replace(r'\text{s}', 's')
        run_text = run_text.replace(r'\text{px}', 'px').replace(r'\text{pixels}', 'pixels').replace(r'\text{fps}', 'fps')
        run_text = run_text.replace(r'\text{KB}', 'KB').replace(r'\text{MB}', 'MB').replace(r'\text{GB}', 'GB')
        run_text = run_text.replace(r'\text{giây}', 'giây').replace(r'\text{E2E}', 'E2E')
        run_text = re.sub(r'\\text\{([^}]+)\}', r'\1', run_text)

        run = p.add_run(run_text)
        f_name = "Consolas" if is_code else base_font
        run.font.name = f_name
        run.font.size = Pt(base_size - 1.0 if is_code else base_size)
        if is_code:
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        else:
            if is_bold:
                run.font.bold = True
            if is_italic:
                run.font.italic = True
            if color:
                run.font.color.rgb = color

        # Font hardening: ensure Western and EastAsia font consistency across all Word environments
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn('w:ascii'), f_name)
        rFonts.set(qn('w:hAnsi'), f_name)
        rFonts.set(qn('w:eastAsia'), f_name)
        rFonts.set(qn('w:cs'), f_name)

def build_docx():
    template_path = 'docs/vi/reports/Báo cáo bài tập nhóm.docx'
    markdown_path = 'docs/vi/reports/BAO_CAO_TONG_HOP.md'
    output_path = 'docs/vi/reports/BAO_CAO_INKSHIFT_FINAL.docx'

    if not os.path.exists(template_path):
        print(f"Template not found: {template_path}")
        return
    if not os.path.exists(markdown_path):
        print(f"Markdown report not found: {markdown_path}")
        return

    print("Loading base document template...")
    doc = docx.Document(template_path)

    # Clean up legacy numbering (numPr) and hanging indents from Heading styles
    # to eliminate double numbering in Word (e.g. '1 PHẦN 1', '1.1 1.1.')
    print("Clearing automatic list numbering and indents from Heading styles...")
    for s in doc.styles:
        if s.name.startswith('Heading') or s.name in ['Subtitle', 'TOC Heading']:
            for numPr in s._element.xpath('.//w:numPr'):
                numPr.getparent().remove(numPr)
            for ind in s._element.xpath('.//w:ind'):
                ind.getparent().remove(ind)

    # 1. Update Cover Page (Table 0)
    print("Updating cover page with InkShift project title...")
    cover_table = doc.tables[0]
    cell = cover_table.rows[0].cells[0]
    
    # In cell, paragraph 7 is the project title: 'Công việc/Dự án: TALKSHOW TALK IT - DO IT 2026'
    # Keep the exact header above (P0 - P5) intact as requested.
    p7 = cell.paragraphs[7]
    p7.text = "" # Clear old text
    r1 = p7.add_run("Công việc/Dự án: ")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(14.0)
    r1.font.bold = True

    r2 = p7.add_run("INKSHIFT — TIỆN ÍCH DỊCH TRUYỆN TRANH THỜI GIAN THỰC ĐA NỀN TẢNG TRÊN TRÌNH DUYỆT")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14.0)
    r2.font.bold = True
    # Keep all other paragraphs (P0-P5 header, P11-P21 members, P27 date) intact as in the original template

    # 2. Clear everything after the cover table
    print("Clearing old report content and slides...")
    body = doc._body._body
    while len(body) > 2:
        body.remove(body[1])

    # 3. Add Page Break after cover table
    doc.add_page_break()

    # 4. Read Markdown and generate content
    print("Reading and parsing BAO_CAO_TONG_HOP.md...")
    with open(markdown_path, 'r', encoding='utf-8') as f:
        md_lines = f.readlines()

    # Filter lines: start from "## PHẦN 1:" (skip intro markdown headers & TOC)
    start_idx = 0
    for i, line in enumerate(md_lines):
        if line.strip().startswith("## PHẦN 1:"):
            start_idx = i
            break

    content_lines = md_lines[start_idx:]
    print(f"Processing {len(content_lines)} markdown lines...")

    in_table = False
    table_lines = []
    in_code_block = False
    code_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines or len(tbl_lines) < 2:
            return
        # Parse rows
        parsed_rows = []
        for l in tbl_lines:
            l = l.strip()
            if not l.startswith('|') or not l.endswith('|'):
                continue
            cells = [c.strip() for c in l[1:-1].split('|')]
            # Check if separator row like |---|---|
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            parsed_rows.append(cells)

        if not parsed_rows:
            return

        num_cols = max(len(r) for r in parsed_rows)
        # Normalize rows to same column count
        for r in parsed_rows:
            while len(r) < num_cols:
                r.append("")

        w_tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        w_tbl.style = 'TableGrid'
        w_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r_idx, row_data in enumerate(parsed_rows):
            w_row = w_tbl.rows[r_idx]
            prevent_row_split(w_row)
            if r_idx == 0:
                set_repeat_header(w_row)

            for c_idx, cell_value in enumerate(row_data):
                cell = w_row.cells[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

                # Style header row
                if r_idx == 0:
                    set_cell_background(cell, "EAECEF") # Light soft gray header
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pf = p.paragraph_format
                    pf.space_after = Pt(2)
                    pf.space_before = Pt(2)
                    pf.line_spacing = 1.15
                    format_inline_markdown(p, cell_value, base_font="Times New Roman", base_size=11.5, is_bold_base=True)
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "FAFAFA") # Subtle zebra striping
                    p = cell.paragraphs[0]
                    pf = p.paragraph_format
                    pf.space_after = Pt(2)
                    pf.space_before = Pt(2)
                    pf.line_spacing = 1.15
                    # Align center if short number/code, else left
                    if len(cell_value) <= 8 and any(char.isdigit() for char in cell_value):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    format_inline_markdown(p, cell_value, base_font="Times New Roman", base_size=11.0)

        # Space after table
        after_p = doc.add_paragraph()
        after_p.paragraph_format.space_before = Pt(0)
        after_p.paragraph_format.space_after = Pt(6)

    def flush_code(c_lines):
        if not c_lines:
            return
        box_text = "\n".join(c_lines)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_before = Pt(4)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.0
        pf.left_indent = Inches(0.2)
        run = p.add_run(box_text)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Process lines
    img_map = {
        "Hình 1:": ("hinh_1_popup_bento_main.png", 3.0, "hinh_1b_popup_settings.png", 3.0),
        "Hình 2:": ("hinh_2_mascot_hud.png", 3.2, None, None),
        "Hình 3:": ("hinh_3a_manga_raw_page11.png", 2.9, "hinh_3b_manga_translated_page11.png", 2.9),
        "Hình 4:": ("04_yolo_segmentation_visualized.png", 4.8, None, None),
        "Hình 5:": ("03_translated_overlay_visualized.png", 4.8, None, None),
        "Hình 6:": ("06_network_waterfall.png", 6.0, None, None),
        "Hình 7:": ("07_circuit_breaker_cache_flow.png", 6.0, None, None),
    }

    current_part_count = 0

    idx = 0
    while idx < len(content_lines):
        line = content_lines[idx].rstrip()
        stripped = line.strip()

        # Handle Code blocks
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code(code_lines)
                code_lines = []
            else:
                if in_table:
                    in_table = False
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
                code_lines = []
            idx += 1
            continue

        if in_code_block:
            code_lines.append(line)
            idx += 1
            continue

        # Handle Tables
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            idx += 1
            continue
        elif in_table:
            in_table = False
            flush_table(table_lines)
            table_lines = []

        # Empty line
        if not stripped:
            idx += 1
            continue

        # Horizontal rule ---
        if stripped == "---":
            idx += 1
            continue

        # Heading 1: ## PHẦN X: ...
        if stripped.startswith("## PHẦN "):
            current_part_count += 1
            # Add page break before Part 2..8
            if current_part_count > 1:
                doc.add_page_break()

            h_text = stripped[2:].strip()
            h = doc.add_heading(level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.LEFT
            hf = h.paragraph_format
            hf.space_before = Pt(18)
            hf.space_after = Pt(12)
            hf.left_indent = Inches(0)
            hf.keep_with_next = True
            for numPr in h._element.xpath('.//w:numPr'):
                numPr.getparent().remove(numPr)
            format_inline_markdown(h, h_text, base_font="Times New Roman", base_size=16.0, is_bold_base=True, color=RGBColor(0x00, 0x00, 0x00))
            idx += 1
            continue

        # Heading 2: ### X.Y. ...
        if stripped.startswith("### "):
            h_text = stripped[3:].strip()
            h = doc.add_heading(level=2)
            h.alignment = WD_ALIGN_PARAGRAPH.LEFT
            hf = h.paragraph_format
            hf.space_before = Pt(18)
            hf.space_after = Pt(6)
            hf.left_indent = Inches(0)
            hf.keep_with_next = True
            for numPr in h._element.xpath('.//w:numPr'):
                numPr.getparent().remove(numPr)
            format_inline_markdown(h, h_text, base_font="Times New Roman", base_size=14.0, is_bold_base=True, color=RGBColor(0x00, 0x00, 0x00))
            idx += 1
            continue

        # Heading 3: #### X.Y.Z. ... or #### Hình X: ...
        if stripped.startswith("#### "):
            h_text = stripped[4:].strip()

            # Check if this is a figure header
            matched_fig = None
            for fig_key in img_map:
                if fig_key in h_text:
                    matched_fig = fig_key
                    break

            h = doc.add_heading(level=3)
            h.alignment = WD_ALIGN_PARAGRAPH.LEFT
            hf = h.paragraph_format
            hf.space_before = Pt(12)
            hf.space_after = Pt(4)
            hf.left_indent = Inches(0)
            hf.keep_with_next = True
            for numPr in h._element.xpath('.//w:numPr'):
                numPr.getparent().remove(numPr)
            format_inline_markdown(h, h_text, base_font="Times New Roman", base_size=13.0, is_bold_base=True, color=RGBColor(0x00, 0x00, 0x00))

            # If it's a figure heading, also insert the image right under it!
            if matched_fig:
                img_info = img_map[matched_fig]
                img1_name, w1, img2_name, w2 = img_info
                img1_path = os.path.join("docs/vi/reports/images", img1_name)

                if os.path.exists(img1_path):
                    img_p = doc.add_paragraph()
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.paragraph_format.space_before = Pt(6)
                    img_p.paragraph_format.space_after = Pt(4)

                    if img2_name:
                        img2_path = os.path.join("docs/vi/reports/images", img2_name)
                        if os.path.exists(img2_path):
                            # Insert two images side by side
                            img_p.add_run().add_picture(img1_path, width=Inches(w1))
                            img_p.add_run("   ") # Small space
                            img_p.add_run().add_picture(img2_path, width=Inches(w2))
                        else:
                            img_p.add_run().add_picture(img1_path, width=Inches(w1 * 1.5))
                    else:
                        img_p.add_run().add_picture(img1_path, width=Inches(w1))

            idx += 1
            continue

        # Check if line is an image caption: *Hình X: ...* or - **Chú thích ảnh báo cáo:** *Hình X: ...*
        caption_match = re.search(r'(\*Hình \d+:.*?\*)', stripped)
        if caption_match and ("Chú thích" in stripped or stripped.startswith("*Hình ")):
            cap_text = caption_match.group(1)
            cap_p = doc.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_before = Pt(4)
            cap_p.paragraph_format.space_after = Pt(12)
            format_inline_markdown(cap_p, cap_text, base_font="Times New Roman", base_size=11.5, is_italic_base=True, color=RGBColor(0x44, 0x44, 0x44))
            idx += 1
            continue

        # Bullet lists (- Item)
        if stripped.startswith("- ") or stripped.startswith("* "):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph(style='Normal')
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(4)
            pf.line_spacing = 1.3
            pf.left_indent = Inches(0.25)
            # Add bullet symbol
            bullet_run = p.add_run("•  ")
            bullet_run.font.name = "Times New Roman"
            bullet_run.font.size = Pt(13.0)
            bullet_run.font.bold = True
            format_inline_markdown(p, item_text, base_font="Times New Roman", base_size=13.0)
            idx += 1
            continue

        # Numbered lists (1. Item, 2. Item)
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            num_str = num_match.group(1)
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='Normal')
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(4)
            pf.line_spacing = 1.3
            pf.left_indent = Inches(0.25)
            prefix_run = p.add_run(f"{num_str}.  ")
            prefix_run.font.name = "Times New Roman"
            prefix_run.font.size = Pt(13.0)
            prefix_run.font.bold = True
            format_inline_markdown(p, item_text, base_font="Times New Roman", base_size=13.0)
            idx += 1
            continue

        # Regular Body Paragraph
        p = doc.add_paragraph(style='Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.5
        format_inline_markdown(p, stripped, base_font="Times New Roman", base_size=13.0)
        idx += 1

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    # 5. Clean up unused legacy slide media relationships from original template
    print("Purging unused legacy slide relationships from template...")
    for rId, rel in list(doc.part.rels.items()):
        if 'image' in rel.reltype:
            fname = rel.target_ref.split('/')[-1]
            if fname.startswith('image') and fname.endswith('.PNG'):
                try:
                    num = int(fname[5:-4])
                    if 3 <= num <= 24:
                        del doc.part.rels[rId]
                except ValueError:
                    pass

    print(f"Saving final cleaned docx to {output_path}...")
    saved_path = output_path
    counter = 1
    while True:
        try:
            doc.save(saved_path)
            break
        except PermissionError:
            counter += 1
            saved_path = f'docs/vi/reports/BAO_CAO_INKSHIFT_FINAL_V{counter}.docx'
            print(f"Warning: Target file is currently locked by Word. Trying {saved_path}...")

    final_size = os.path.getsize(saved_path)
    print(f"Successfully generated: {saved_path} ({final_size / 1024 / 1024:.2f} MB)")

if __name__ == '__main__':
    build_docx()

The above content shows the entire, complete file contents of the requested file.